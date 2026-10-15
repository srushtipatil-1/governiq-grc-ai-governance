"""Phase 5 tests: coverage, gaps, crosswalk and the Risk & Control Matrix."""

from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from governiq.controls import (
    RCM_COLUMNS,
    SAMPLE_SIZES,
    control_rows,
    crosswalk,
    framework_coverage,
    frameworks_of,
    load_controls,
    rcm_rows,
    rcm_workbook,
    risks_by_control,
)
from governiq.models import Control, Risk


def _coverage(session) -> dict:
    return {cov.framework_code: cov for cov in framework_coverage(session)}


def test_coverage_counts_mapped_requirements(session):
    cov = _coverage(session)
    assert (cov["ISO27001"].mapped, cov["ISO27001"].total) == (17, 20)
    assert cov["ISO27001"].percent == 85
    assert (cov["DPDP"].mapped, cov["DPDP"].total) == (8, 8)


def test_gaps_name_the_unmapped_requirements(session):
    cov = _coverage(session)
    assert [g.ref for g in cov["ISO27001"].gaps] == ["A.5.23", "A.5.30", "A.8.16"]
    assert cov["DPDP"].gaps == []
    assert sum(len(c.gaps) for c in cov.values()) == 13


def test_adding_a_mapping_closes_a_gap(session):
    """Mapping a control to a gap requirement must remove it from the gap list."""
    cloud_gap = _coverage(session)["ISO27001"].gaps[0]
    control = session.scalar(select(Control).where(Control.code == "CTL-TPR-01"))
    control.requirements.append(cloud_gap)
    session.commit()
    assert "A.5.23" not in [g.ref for g in _coverage(session)["ISO27001"].gaps]


def test_access_review_counts_for_three_frameworks(session):
    review = next(c for c in load_controls(session) if c.code == "CTL-IAM-03")
    assert frameworks_of(review) == ["DPDP", "ISO 27001", "NIST CSF"]


def test_control_rows_report_linked_risks(session):
    controls = load_controls(session)
    rows = {row["Code"]: row for row in control_rows(controls, risks_by_control(session))}
    assert rows["CTL-AI-09"]["Risks"] == 3  # RSK-023, RSK-024, RSK-025
    assert rows["CTL-IAM-03"]["Requirements"] == 4


def test_crosswalk_totals_match_all_mappings(session):
    controls = load_controls(session)
    grid = crosswalk(controls)
    assert sum(sum(row.values()) for row in grid.values()) == sum(len(c.requirements) for c in controls)


def test_rcm_has_one_row_per_risk_control_pair(session):
    pairs = sum(len(r.controls) for r in session.scalars(select(Risk)))
    rows = rcm_rows(session)
    assert len(rows) == pairs == 47
    assert list(rows[0]) == RCM_COLUMNS
    assert {row["Test result"] for row in rows} == {"Not tested"}


@pytest.mark.parametrize(("frequency", "size"), [("Annual", "1"), ("Quarterly", "2"), ("Daily", "20")])
def test_sample_size_follows_control_frequency(session, frequency, size):
    row = next(r for r in rcm_rows(session) if r["Frequency"] == frequency)
    assert row["Sample size"] == size == SAMPLE_SIZES[frequency]


def test_rcm_workbook_opens_with_expected_sheets(session):
    book = load_workbook(BytesIO(rcm_workbook(session)))
    assert book.sheetnames == ["Risk & Control Matrix", "Framework coverage", "Read me"]
    rcm = book["Risk & Control Matrix"]
    assert [cell.value for cell in rcm[1]] == RCM_COLUMNS
    assert rcm.max_row == 48  # header + 47 rows
    assert "Not an audit opinion" in book["Read me"]["A2"].value
