"""Control library logic: framework coverage, gaps, crosswalk and the Risk & Control Matrix (RCM).

Coverage answers "does at least one control map to this requirement?". It does not say the control
works: effectiveness needs control tests, which arrive in Phase 6.
"""

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from governiq.config import FRAMEWORK_LABELS
from governiq.models import Control, Framework, Requirement, Risk

# Illustrative sample sizes for operating-effectiveness tests, by how often a control runs.
# Real firms set their own table; this follows the common pattern used in audit practice.
SAMPLE_SIZES = {
    "Annual": "1",
    "Quarterly": "2",
    "Monthly": "2",
    "Weekly": "5",
    "Daily": "20",
    "Per event": "25 (or all, if fewer occurred)",
    "Continuous": "1 (automated; rely on IT general controls)",
}


def load_controls(session: Session) -> list[Control]:
    query = (
        select(Control)
        .options(
            selectinload(Control.owner),
            selectinload(Control.requirements).selectinload(Requirement.framework),
        )
        .order_by(Control.code)
    )
    return list(session.scalars(query))


def risks_by_control(session: Session) -> dict[str, list[Risk]]:
    """Which risks each control treats, keyed by control code."""
    linked: dict[str, list[Risk]] = {}
    risks = session.scalars(select(Risk).options(selectinload(Risk.controls)).order_by(Risk.code))
    for risk in risks:
        for control in risk.controls:
            linked.setdefault(control.code, []).append(risk)
    return linked


def frameworks_of(control: Control) -> list[str]:
    codes = sorted({req.framework.code for req in control.requirements})
    return [FRAMEWORK_LABELS[code] for code in codes]


def control_rows(controls: list[Control], linked: dict[str, list[Risk]]) -> list[dict]:
    return [
        {
            "Code": c.code,
            "Control": c.title,
            "Domain": c.domain,
            "Owner": c.owner.full_name,
            "Type": c.control_type,
            "Nature": c.nature,
            "Frequency": c.frequency,
            "Key": "Yes" if c.is_key else "No",
            "Frameworks": ", ".join(frameworks_of(c)),
            "Requirements": len(c.requirements),
            "Risks": len(linked.get(c.code, [])),
        }
        for c in controls
    ]


@dataclass
class Coverage:
    framework_code: str
    name: str
    claim_level: str
    total: int
    mapped: int
    gaps: list[Requirement]

    @property
    def label(self) -> str:
        return FRAMEWORK_LABELS[self.framework_code]

    @property
    def percent(self) -> int:
        return round(100 * self.mapped / self.total) if self.total else 0


def framework_coverage(session: Session) -> list[Coverage]:
    """For each framework: how many requirements have at least one mapped control, and which do not."""
    mapped_ids = {req.id for control in load_controls(session) for req in control.requirements}
    frameworks = session.scalars(
        select(Framework).options(selectinload(Framework.requirements)).order_by(Framework.id)
    )
    result = []
    for fw in frameworks:
        reqs = sorted(fw.requirements, key=lambda r: r.id)
        gaps = [r for r in reqs if r.id not in mapped_ids]
        result.append(Coverage(fw.code, fw.name, fw.claim_level, len(reqs), len(reqs) - len(gaps), gaps))
    return result


def crosswalk(controls: list[Control]) -> dict[str, dict[str, int]]:
    """Requirements covered per control domain and framework: {domain: {framework label: count}}."""
    table: dict[str, dict[str, int]] = {}
    for control in controls:
        row = table.setdefault(control.domain, dict.fromkeys(FRAMEWORK_LABELS.values(), 0))
        for req in control.requirements:
            row[FRAMEWORK_LABELS[req.framework.code]] += 1
    return table


RCM_COLUMNS = [
    "Risk ID",
    "Risk",
    "Control ID",
    "Control",
    "Control owner",
    "Key control",
    "Type",
    "Nature",
    "Frequency",
    "Framework references",
    "Test of design",
    "Test of operating effectiveness",
    "Sample size",
    "Test result",
]


def rcm_rows(session: Session) -> list[dict]:
    """One row per risk-control pair, in the layout auditors use for a Risk & Control Matrix."""
    rows = []
    linked = risks_by_control(session)
    for control in load_controls(session):
        for risk in linked.get(control.code, []):
            refs = "; ".join(f"{FRAMEWORK_LABELS[r.framework.code]} {r.ref}" for r in control.requirements)
            sample = SAMPLE_SIZES.get(control.frequency, "Set by the engagement team")
            rows.append(
                {
                    "Risk ID": risk.code,
                    "Risk": risk.title,
                    "Control ID": control.code,
                    "Control": control.title,
                    "Control owner": control.owner.full_name,
                    "Key control": "Yes" if control.is_key else "No",
                    "Type": control.control_type,
                    "Nature": control.nature,
                    "Frequency": control.frequency,
                    "Framework references": refs,
                    "Test of design": (
                        f"Inquire with {control.owner.full_name} and inspect the documented procedure to "
                        f"confirm the control is designed to address {risk.code}."
                    ),
                    "Test of operating effectiveness": (
                        f"Select a sample of {control.frequency.lower()} occurrences and inspect evidence "
                        "that the control operated as designed throughout the period."
                    ),
                    "Sample size": sample,
                    "Test result": "Not tested",
                }
            )
    return sorted(rows, key=lambda row: (row["Risk ID"], row["Control ID"]))


def rcm_workbook(session: Session) -> bytes:
    """Build the RCM and coverage summary as a formatted Excel workbook."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Risk & Control Matrix"
    header_fill = PatternFill("solid", fgColor="14213D")
    header_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")

    sheet.append(RCM_COLUMNS)
    for row in rcm_rows(session):
        sheet.append([row[col] for col in RCM_COLUMNS])
    widths = [10, 40, 12, 40, 18, 10, 12, 11, 12, 40, 45, 45, 22, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill, cell.font, cell.alignment = header_fill, header_font, wrap
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = wrap
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions

    summary = book.create_sheet("Framework coverage")
    summary.append(
        ["Framework", "Claim level", "Requirements", "Mapped", "Coverage %", "Unmapped requirements"]
    )
    for cov in framework_coverage(session):
        gaps = "; ".join(f"{g.ref} {g.title}" for g in cov.gaps) or "None"
        summary.append([cov.name, cov.claim_level, cov.total, cov.mapped, cov.percent, gaps])
    for index, width in enumerate([42, 20, 14, 10, 12, 80], start=1):
        summary.column_dimensions[get_column_letter(index)].width = width
    for cell in summary[1]:
        cell.fill, cell.font, cell.alignment = header_fill, header_font, wrap
    for line in summary.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = wrap

    notes = book.create_sheet("Read me")
    for line in (
        ["GovernIQ Risk & Control Matrix (synthetic data, Kshitij Finserv Pvt Ltd, fictional)"],
        ["Educational portfolio output. Not an audit opinion, certification or conformity assessment."],
        ["Coverage means at least one control maps to a requirement; it does not mean the control works."],
        ["Sample sizes are illustrative; real engagements use their firm's own sampling methodology."],
    ):
        notes.append(line)
    notes.column_dimensions["A"].width = 100

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
